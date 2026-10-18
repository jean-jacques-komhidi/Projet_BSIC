# -*- coding: utf-8 -*-
"""
Export des donnees pour Power BI, Excel ou une autorite de controle.

Trois jeux de donnees : les analyses de credit, les clients et les versions du
modele. Les colonnes sont regroupees ; l'administrateur choisit les groupes a
transmettre. Les noms des clients peuvent etre remplaces par un code, et le
genre n'est exporte que s'il est explicitement demande.

Chaque fichier est accompagne d'un dictionnaire des donnees (signification,
unite et valeurs possibles de chaque colonne). Les mots de passe et les
conversations de l'assistant ne sont jamais exportes.
"""

import io
import json
import csv
import zipfile
from datetime import datetime, timedelta

from sqlalchemy.orm import Session

from . import models
from .pdf_service import CATEGORIES

DECALAGE_TCHAD = timedelta(hours=1)
JEUX = ("analyses", "clients", "versions")
DESTINATAIRES = ("Analyse interne (Power BI)", "Autorité de contrôle (COBAC)", "Audit", "Autre")

# Groupes de colonnes proposes a l'administrateur
GROUPES = {
    "identification": "Identification",
    "financier": "Situation financière",
    "modele": "Résultat du modèle",
    "issue": "Issue réelle",
    "profil": "Profil du client",
    "genre": "Genre",
}
GROUPES_PAR_DEFAUT = ("identification", "financier", "modele", "issue")

# Colonnes : (nom, groupe, signification, unite, valeurs possibles)
COLONNES = {
    "analyses": [
        ("reference", "identification", "Référence de l'analyse", "", "CRED-AAAA-NNNN"),
        ("date_analyse", "identification", "Date et heure de l'analyse (heure du Tchad)", "AAAA-MM-JJ HH:MM", ""),
        ("agent", "identification", "Agent de crédit qui a réalisé l'analyse", "", ""),
        ("client", "identification", "Client (nom, ou code si les noms sont anonymisés)", "", "CLI-NNNN si anonymisé"),
        ("montant_fcfa", "financier", "Montant du crédit demandé", "FCFA", ""),
        ("revenu_fcfa", "financier", "Revenu du demandeur", "FCFA", ""),
        ("echeance_fcfa", "financier", "Montant de l'échéance", "FCFA", ""),
        ("taux_endettement", "financier", "Échéance divisée par le revenu", "ratio (0 à 1)", ""),
        ("anciennete_emploi_ans", "financier", "Ancienneté dans l'emploi", "années", ""),
        ("probabilite_defaut", "modele", "Probabilité de défaut calculée par le modèle", "ratio (0 à 1)", ""),
        ("classe_risque", "modele", "Classe de risque", "", "faible, moyen, eleve"),
        ("decision", "modele", "Décision proposée", "", "ACCORDE, REFUSE"),
        ("version_modele", "modele", "Version du modèle qui a noté le dossier", "", "v1, v2…"),
        ("issue_reelle", "issue", "Issue réelle du crédit", "",
         "rembourse, defaut, en attente (accordé, issue pas encore connue), vide si refusé"),
        ("age_ans", "profil", "Âge du demandeur", "années", ""),
        ("profession", "profil", "Profession du client", "", ""),
        ("situation_familiale", "profil", "Situation familiale du client", "", ""),
        ("genre", "genre", "Genre du client (donnée sensible)", "", "F, M"),
    ],
    "clients": [
        ("client", "identification", "Client (nom, ou code si les noms sont anonymisés)", "", "CLI-NNNN si anonymisé"),
        ("date_creation", "identification", "Date d'enregistrement du client", "AAAA-MM-JJ", ""),
        ("nb_analyses", "identification", "Nombre d'analyses sur la période", "", ""),
        ("derniere_analyse", "identification", "Date de la dernière analyse sur la période", "AAAA-MM-JJ", ""),
        ("derniere_probabilite", "modele", "Probabilité de défaut de la dernière analyse", "ratio (0 à 1)", ""),
        ("derniere_decision", "modele", "Décision de la dernière analyse", "", "ACCORDE, REFUSE"),
        ("credits_rembourses", "issue", "Crédits remboursés sur la période", "", ""),
        ("credits_en_defaut", "issue", "Crédits en défaut sur la période", "", ""),
        ("age_ans", "profil", "Âge du client", "années", ""),
        ("profession", "profil", "Profession du client", "", ""),
        ("situation_familiale", "profil", "Situation familiale du client", "", ""),
        ("genre", "genre", "Genre du client (donnée sensible)", "", "F, M"),
    ],
    "versions": [
        ("version", "identification", "Numéro de version du modèle", "", "v1, v2…"),
        ("algorithme", "identification", "Algorithme d'apprentissage", "", ""),
        ("statut", "identification", "Statut de la version", "", "production, retire, rejete"),
        ("date_mise_en_service", "identification", "Date de création ou de mise en service", "AAAA-MM-JJ", ""),
        ("date_retrait", "identification", "Date de retrait", "AAAA-MM-JJ", ""),
        ("auc_entrainement", "modele", "AUC sur le jeu de test du socle d'entraînement", "0,5 à 1", ""),
        ("auc_reelle", "issue", "AUC mesurée sur des issues réelles mises de côté", "0,5 à 1", ""),
        ("nb_issues_validation", "issue", "Nombre d'issues réelles utilisées pour l'AUC réelle", "", ""),
        ("auc_modele_precedent", "issue", "AUC du modèle précédent sur les mêmes issues", "0,5 à 1", ""),
    ],
}


def _locale(d):
    return d + DECALAGE_TCHAD if d else None


def _date(d, heure=False):
    d = _locale(d)
    return d.strftime("%Y-%m-%d %H:%M" if heure else "%Y-%m-%d") if d else ""


def code_client(client_id) -> str:
    return f"CLI-{client_id:04d}" if client_id else ""


def _dossier(a) -> dict:
    try:
        return json.loads(a.donnees_dossier or "{}")
    except (TypeError, ValueError):
        return {}


def _fr(valeur) -> str:
    """Libelle en francais (situation familiale, etc.)."""
    return CATEGORIES.get(valeur, valeur) if valeur else ""


def _arrondi(x, n=4):
    return round(float(x), n) if x is not None else None


def colonnes(jeu: str, groupes) -> list:
    """Colonnes retenues, dans l'ordre, pour les groupes choisis."""
    return [c[0] for c in COLONNES[jeu] if c[1] in groupes]


def _periode(du, au):
    """Bornes en temps universel ; 'au' inclut toute la journee."""
    debut = datetime.combine(du, datetime.min.time()) - DECALAGE_TCHAD if du else None
    fin = datetime.combine(au, datetime.max.time()) - DECALAGE_TCHAD if au else None
    return debut, fin


def _analyses(db: Session, debut, fin):
    q = db.query(models.Analyse)
    if debut:
        q = q.filter(models.Analyse.date_analyse >= debut)
    if fin:
        q = q.filter(models.Analyse.date_analyse <= fin)
    return q.order_by(models.Analyse.date_analyse.desc()).all()


def _numeros_versions(db: Session) -> dict:
    ids = [v.id for v in db.query(models.ModelVersion).order_by(models.ModelVersion.id).all()]
    return {vid: f"v{i}" for i, vid in enumerate(ids, start=1)}


def lignes(db: Session, jeu: str, groupes, anonymiser=True, du=None, au=None) -> list:
    """Lignes du jeu demande (dictionnaires limites aux colonnes choisies)."""
    debut, fin = _periode(du, au)
    retenues = colonnes(jeu, groupes)
    nom_client = (lambda c: code_client(c.id)) if anonymiser else (lambda c: c.nom)
    resultat = []

    if jeu == "analyses":
        users = {u.id: u.nom for u in db.query(models.User).all()}
        clients = {c.id: c for c in db.query(models.Client).all()}
        versions = _numeros_versions(db)
        for a in _analyses(db, debut, fin):
            d = _dossier(a)
            c = clients.get(a.client_id)
            revenu, echeance = d.get("AMT_INCOME_TOTAL"), d.get("AMT_ANNUITY")
            if a.resultat_reel in ("rembourse", "defaut"):
                issue = a.resultat_reel
            else:
                issue = "en attente" if a.decision == "ACCORDE" else ""
            date_ref = a.date_analyse or datetime.utcnow()
            ligne = {
                "reference": f"CRED-{date_ref.year}-{a.id:04d}",
                "date_analyse": _date(a.date_analyse, heure=True),
                "agent": users.get(a.user_id, ""),
                "client": nom_client(c) if c else "",
                "montant_fcfa": d.get("AMT_CREDIT"),
                "revenu_fcfa": revenu,
                "echeance_fcfa": echeance,
                "taux_endettement": _arrondi(echeance / revenu, 3) if revenu and echeance else None,
                "anciennete_emploi_ans": d.get("ANCIENNETE_EMPLOI_ANNEES"),
                "probabilite_defaut": _arrondi(a.probabilite_defaut),
                "classe_risque": a.classe_risque or "",
                "decision": a.decision or "",
                "version_modele": versions.get(a.model_version_id, ""),
                "issue_reelle": issue,
                "age_ans": (c.age if c and c.age else d.get("AGE_ANNEES")),
                "profession": c.profession if c and c.profession else "",
                "situation_familiale": _fr(c.situation_familiale) if c else "",
                "genre": (c.genre if c and c.genre else d.get("CODE_GENDER", "")),
            }
            resultat.append({k: ligne[k] for k in retenues})

    elif jeu == "clients":
        par_client = {}
        for a in _analyses(db, debut, fin):
            par_client.setdefault(a.client_id, []).append(a)
        for c in db.query(models.Client).order_by(models.Client.id).all():
            liste = par_client.get(c.id)
            if not liste:
                continue
            derniere = liste[0]   # les analyses sont triees de la plus recente a la plus ancienne
            ligne = {
                "client": nom_client(c),
                "date_creation": _date(c.date_creation),
                "nb_analyses": len(liste),
                "derniere_analyse": _date(derniere.date_analyse),
                "derniere_probabilite": _arrondi(derniere.probabilite_defaut),
                "derniere_decision": derniere.decision or "",
                "credits_rembourses": sum(a.resultat_reel == "rembourse" for a in liste),
                "credits_en_defaut": sum(a.resultat_reel == "defaut" for a in liste),
                "age_ans": c.age,
                "profession": c.profession or "",
                "situation_familiale": _fr(c.situation_familiale),
                "genre": c.genre or "",
            }
            resultat.append({k: ligne[k] for k in retenues})

    else:  # versions du modele (toutes, la periode ne s'applique pas)
        for i, v in enumerate(db.query(models.ModelVersion).order_by(models.ModelVersion.id).all(), start=1):
            ligne = {
                "version": f"v{i}",
                "algorithme": v.algorithme,
                "statut": "production" if v.en_production else (v.statut or "retire"),
                "date_mise_en_service": _date(v.date_creation),
                "date_retrait": _date(v.date_retrait),
                "auc_entrainement": _arrondi(v.auc),
                "auc_reelle": _arrondi(v.auc_reelle),
                "nb_issues_validation": v.nb_issues_validation,
                "auc_modele_precedent": _arrondi(v.auc_precedent),
            }
            resultat.append({k: ligne[k] for k in retenues})
        resultat.reverse()
    return resultat


def dictionnaire(jeu: str, groupes) -> list:
    return [{"colonne": nom, "groupe": GROUPES[g], "signification": sens, "unite": unite, "valeurs_possibles": valeurs}
            for nom, g, sens, unite, valeurs in COLONNES[jeu] if g in groupes]


# ---------------------------------------------------------------------------
# Fichiers
# ---------------------------------------------------------------------------
def _valeur_csv(v):
    """Nombres a virgule decimale (Excel et Power BI en francais), vide pour None."""
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v}".replace(".", ",")
    return v


def _csv(lignes_, entetes) -> bytes:
    tampon = io.StringIO()
    ecrivain = csv.writer(tampon, delimiter=";", lineterminator="\r\n")
    ecrivain.writerow(entetes)
    for ligne in lignes_:
        ecrivain.writerow([_valeur_csv(ligne.get(k)) for k in entetes])
    # BOM UTF-8 : Excel reconnait les accents a l'ouverture
    return ("﻿" + tampon.getvalue()).encode("utf-8")


def fichier_zip_csv(jeu, lignes_, entetes, dico) -> bytes:
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{jeu}.csv", _csv(lignes_, entetes))
        z.writestr("dictionnaire_des_donnees.csv",
                   _csv(dico, ["colonne", "groupe", "signification", "unite", "valeurs_possibles"]))
    return tampon.getvalue()


def fichier_xlsx(jeu, lignes_, entetes, dico) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Données"
    for titre, rangs, cles in ((feuille, lignes_, entetes),
                               (classeur.create_sheet("Dictionnaire"), dico,
                                ["colonne", "groupe", "signification", "unite", "valeurs_possibles"])):
        titre.append(cles)
        for r in rangs:
            titre.append([r.get(k) for k in cles])
        for cellule in titre[1]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="0055A8")
        titre.freeze_panes = "A2"
        for i, k in enumerate(cles, start=1):
            largeur = max([len(str(k))] + [len(str(r.get(k) or "")) for r in rangs[:200]])
            titre.column_dimensions[get_column_letter(i)].width = min(max(largeur + 2, 10), 60)
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
