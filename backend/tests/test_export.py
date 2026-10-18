# -*- coding: utf-8 -*-
"""
Tests de l'export des donnees : reserve a l'administrateur, anonymisation,
genre exclu par defaut, destinataire et motif obligatoires, journal, fichiers
CSV (zip avec dictionnaire) et Excel (deux feuilles).
"""

import io
import os
import sys
import json
import zipfile
import tempfile
from datetime import datetime, timedelta
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

_base = os.path.join(tempfile.mkdtemp(), "test_export.db")
os.environ["DATABASE_URL"] = f"sqlite:///{_base}"

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from app.database import Base, engine, SessionLocal
from app import models, export_routes
from app.security import administrateur_courant

app = FastAPI()
app.include_router(export_routes.router)
client = TestClient(app)
DOSSIER = {"AMT_CREDIT": 2_000_000, "AMT_ANNUITY": 100_000, "AMT_INCOME_TOTAL": 250_000,
           "ANCIENNETE_EMPLOI_ANNEES": 4.5, "AGE_ANNEES": 35, "CODE_GENDER": "F"}
DEMANDE = {"jeu": "analyses", "destinataire": "Analyse interne (Power BI)",
           "motif": "Tableau de bord mensuel du risque", "format": "csv"}


@pytest.fixture(autouse=True)
def base():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    admin = models.User(nom="Admin", email="a@x.td", mot_de_passe="secret-hache", role="admin")
    agent = models.User(nom="Fatimé Abakar", email="g@x.td", mot_de_passe="x", role="agent")
    c = models.Client(nom="Achta Oumar", genre="F", age=35, profession="Commerçante", situation_familiale="Married")
    v = models.ModelVersion(algorithme="Gradient Boosting", auc=0.753, en_production=True)
    db.add_all([admin, agent, c, v])
    db.commit()
    for jours, decision, issue in [(200, "ACCORDE", "defaut"), (90, "ACCORDE", "rembourse"),
                                   (40, "ACCORDE", None), (2, "REFUSE", None)]:
        db.add(models.Analyse(client_id=c.id, user_id=agent.id, decision=decision, probabilite_defaut=0.4,
                              classe_risque="moyen", donnees_dossier=json.dumps(DOSSIER), resultat_reel=issue,
                              model_version_id=v.id, date_analyse=datetime.utcnow() - timedelta(days=jours)))
    db.commit()
    utilisateur = SimpleNamespace(id=admin.id, role="admin", nom="Admin", email="a@x.td")
    db.close()
    app.dependency_overrides[administrateur_courant] = lambda: utilisateur
    yield
    app.dependency_overrides.clear()


def _csv(reponse, nom="analyses.csv"):
    z = zipfile.ZipFile(io.BytesIO(reponse.content))
    return z.read(nom).decode("utf-8-sig").splitlines()


def test_reserve_a_l_administrateur():
    def refuse():
        raise HTTPException(status_code=403, detail="Accès réservé à l'administrateur")
    app.dependency_overrides[administrateur_courant] = refuse
    assert client.get("/export/journal").status_code == 403
    assert client.post("/export/telecharger", json=DEMANDE).status_code == 403


def test_noms_anonymises_et_genre_exclu_par_defaut():
    lignes = _csv(client.post("/export/telecharger", json=DEMANDE))
    entetes = lignes[0].split(";")
    assert "genre" not in entetes and "profession" not in entetes
    assert "Achta" not in "\n".join(lignes)
    assert lignes[1].split(";")[entetes.index("client")] == "CLI-0001"


def test_genre_et_noms_seulement_si_demandes():
    demande = {**DEMANDE, "anonymiser": False,
               "groupes": ["identification", "profil", "genre"]}
    lignes = _csv(client.post("/export/telecharger", json=demande))
    assert "genre" in lignes[0].split(";") and "Achta Oumar" in lignes[1]
    assert "Marié(e)" in lignes[1]


def test_destinataire_et_motif_obligatoires():
    assert client.post("/export/telecharger", json={**DEMANDE, "motif": " "}).status_code == 422
    assert client.post("/export/telecharger", json={**DEMANDE, "destinataire": "Quelqu'un"}).status_code == 422
    assert client.post("/export/telecharger", json={**DEMANDE, "groupes": []}).status_code == 422


def test_chaque_export_est_journalise():
    client.post("/export/telecharger", json={**DEMANDE, "destinataire": "Autorité de contrôle (COBAC)",
                                            "motif": "Contrôle du modèle de notation"})
    journal = client.get("/export/journal").json()
    assert len(journal) == 1
    e = journal[0]
    assert e["destinataire"] == "Autorité de contrôle (COBAC)" and e["motif"] == "Contrôle du modèle de notation"
    assert e["anonymise"] and e["nb_lignes"] == 4 and e["auteur"] == "Admin" and not e["genre_inclus"]


def test_periode_et_issue_reelle():
    debut = (datetime.utcnow() - timedelta(days=100)).date().isoformat()
    lignes = _csv(client.post("/export/telecharger", json={**DEMANDE, "du": debut}))
    assert len(lignes) == 4          # en-tete + 3 analyses des 100 derniers jours
    entetes = lignes[0].split(";")
    issues = [l.split(";")[entetes.index("issue_reelle")] for l in lignes[1:]]
    assert issues == ["", "en attente", "rembourse"]


def test_dictionnaire_joint_et_excel_a_deux_feuilles():
    dico = _csv(client.post("/export/telecharger", json=DEMANDE), "dictionnaire_des_donnees.csv")
    assert dico[0].startswith("colonne;groupe;signification")
    assert any(l.startswith("montant_fcfa;") and "FCFA" in l for l in dico)
    from openpyxl import load_workbook
    r = client.post("/export/telecharger", json={**DEMANDE, "format": "xlsx"})
    classeur = load_workbook(io.BytesIO(r.content))
    assert classeur.sheetnames == ["Données", "Dictionnaire"]
    assert classeur["Données"].max_row == 5


def test_apercu_et_jamais_de_mot_de_passe():
    a = client.post("/export/apercu", json={"jeu": "analyses"}).json()
    assert a["total"] == 4 and len(a["lignes"]) == 4
    for jeu in ("clients", "versions"):
        r = client.post("/export/telecharger", json={**DEMANDE, "jeu": jeu, "groupes": list(export_routes.ex.GROUPES)})
        assert b"secret-hache" not in zipfile.ZipFile(io.BytesIO(r.content)).read(f"{jeu}.csv")
